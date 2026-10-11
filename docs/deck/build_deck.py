#!/usr/bin/env python3
"""Generate docs/deck/ScamShield-Immune.pptx (16:9, dark theme) with python-pptx.

Run from anywhere:   python3 docs/deck/build_deck.py
Needs: python-pptx, Pillow. Reads docs/screenshots/*.png and the benchmark JSON in
services/demo_server/demo_server/data/benchmark_summary.json. Every number on a chart comes from
that JSON at build time; the other figures are quoted from the repo docs named in each slide's
speaker notes. Nothing is invented. Render check: see docs/deck/README.md.
"""
from __future__ import annotations

import io
import json
import pathlib
import sys

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Emu, Inches, Pt
from lxml import etree

ROOT = pathlib.Path(__file__).resolve().parents[2]
SHOTS = ROOT / "docs" / "screenshots"
BENCH = json.loads((ROOT / "services/demo_server/demo_server/data/benchmark_summary.json").read_text())
OUT = ROOT / "docs" / "deck" / "ScamShield-Immune.pptx"

# ---- design tokens (match the console) -------------------------------------------------------
INK, CARD, LINE = "0B0F17", "151B28", "2A3446"
TEXT, MUTED, SUBTLE = "E8ECF4", "9FB0C8", "6F819B"
EMERALD, AMBER, CRIMSON, VIOLET, CYAN = "34D399", "F59E0B", "F43F5E", "A78BFA", "22D3EE"
FONT = "Arial"
W, H = 13.333, 7.5
MX = 0.6  # side margin
TOTAL = 14


def rgb(h: str) -> RGBColor:
    return RGBColor.from_string(h)


prs = Presentation()
prs.slide_width, prs.slide_height = Inches(W), Inches(H)
BLANK = prs.slide_layouts[6]
WORDS: dict[int, int] = {}


class S:
    """One slide plus helpers; keeps a body word count (titles, kickers, footers excluded)."""

    def __init__(self, kicker: str, title: str | None, data: bool = True):
        self.s = prs.slides.add_slide(BLANK)
        self.n = len(prs.slides)
        WORDS[self.n] = 0
        bg = self.s.background.fill
        bg.solid()
        bg.fore_color.rgb = rgb(INK)
        if title is not None:
            self.text(MX, 0.42, 9, 0.3, kicker.upper(), 11, EMERALD, bold=True, count=False, spacing=150)
            self.text(MX, 0.72, W - 2 * MX, 0.75, title, 30, TEXT, bold=True, count=False, anchor="top")
        self.text(MX, 7.02, 4, 0.25, "Simulation data" if data else "", 10, SUBTLE, count=False)
        self.text(W - MX - 4, 7.02, 4, 0.25, f"ScamShield Immune   {self.n} / {TOTAL}", 10, SUBTLE, align="r", count=False)

    # -- primitives
    def text(self, x, y, w, h, content, size=14, color=TEXT, bold=False, align="l", anchor="top",
             italic=False, count=True, spacing=0, mono=False, line=None, after=0):
        tb = self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = {"top": MSO_ANCHOR.TOP, "mid": MSO_ANCHOR.MIDDLE, "bot": MSO_ANCHOR.BOTTOM}[anchor]
        paras = content if isinstance(content, list) else [content]
        for i, p in enumerate(paras):
            para = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            opts = {}
            if isinstance(p, tuple):
                p, opts = p
            para.alignment = {"l": PP_ALIGN.LEFT, "c": PP_ALIGN.CENTER, "r": PP_ALIGN.RIGHT}[align]
            if line:
                para.line_spacing = line
            para.space_after = Pt(opts.get("after", after))
            r = para.add_run()
            r.text = p
            f = r.font
            f.name = "Courier New" if (mono or opts.get("mono")) else FONT
            f.size = Pt(opts.get("size", size))
            f.bold = opts.get("bold", bold)
            f.italic = italic
            f.color.rgb = rgb(opts.get("color", color))
            if spacing:
                r._r.get_or_add_rPr().set("spc", str(spacing))
            if count:
                WORDS[self.n] += len(p.split())
        return tb

    def card(self, x, y, w, h, fill=CARD, line=LINE, radius=0.05, alpha=None, lw=0.75):
        sh = self.s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(x), Inches(y), Inches(w), Inches(h))
        sh.adjustments[0] = radius
        sh.fill.solid()
        sh.fill.fore_color.rgb = rgb(fill)
        if alpha is not None:
            clr = sh.fill._xPr.find(qn("a:solidFill")).find(qn("a:srgbClr"))
            a = etree.SubElement(clr, qn("a:alpha"))
            a.set("val", str(int(alpha * 1000)))
        if line:
            sh.line.color.rgb = rgb(line)
            sh.line.width = Pt(lw)
        else:
            sh.line.fill.background()
        sh.shadow.inherit = False
        return sh

    def chip(self, x, y, w, h, label, color=EMERALD, size=11, fill=None):
        sh = self.card(x, y, w, h, fill=fill or INK, line=color, radius=0.5, lw=1)
        tf = sh.text_frame
        tf.margin_left = tf.margin_right = Inches(0.06)
        tf.margin_top = tf.margin_bottom = 0
        tf.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = tf.paragraphs[0]
        p.alignment = PP_ALIGN.CENTER
        r = p.add_run()
        r.text = label
        r.font.size, r.font.bold, r.font.name = Pt(size), True, FONT
        r.font.color.rgb = rgb(color)
        WORDS[self.n] += len(label.split())
        return sh

    def arrow(self, x1, y1, x2, y2, color=MUTED, w=1.5):
        c = self.s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(y1), Inches(x2), Inches(y2))
        c.line.color.rgb = rgb(color)
        c.line.width = Pt(w)
        ln = c.line._get_or_add_ln()
        tail = etree.SubElement(ln, qn("a:tailEnd"))
        tail.set("type", "triangle")
        return c

    def image(self, name: str, x, y, w, h, crop_h: int | None = None, border=True, pad=0.0):
        """Place a screenshot inside box (x, y, w, h), preserving aspect, centred."""
        im = Image.open(SHOTS / name)
        if crop_h:
            im = im.crop((0, 0, im.width, crop_h))
        ar = im.width / im.height
        bw, bh = w - 2 * pad, h - 2 * pad
        if bw / bh > ar:
            ph, pw = bh, bh * ar
        else:
            pw, ph = bw, bw / ar
        px, py = x + (w - pw) / 2, y + (h - ph) / 2
        if border:
            self.card(px - 0.06, py - 0.06, pw + 0.12, ph + 0.12, fill=CARD, line=LINE, radius=0.02)
        buf = io.BytesIO()
        im.save(buf, "PNG")
        buf.seek(0)
        self.s.shapes.add_picture(buf, Inches(px), Inches(py), Inches(pw), Inches(ph))
        return px, py, pw, ph

    def stat(self, x, y, w, h, value, label, sub=None, color=EMERALD, vsize=34):
        self.card(x, y, w, h)
        self.text(x + 0.2, y + 0.14, w - 0.4, 0.7, value, vsize, color, bold=True)
        self.text(x + 0.2, y + 0.14 + vsize / 72 * 1.25, w - 0.4, 0.3, label, 12, TEXT, bold=True)
        if sub:
            self.text(x + 0.2, y + 0.14 + vsize / 72 * 1.25 + 0.3, w - 0.4, 0.3, sub, 11, MUTED)

    def notes(self, text: str):
        self.s.notes_slide.notes_text_frame.text = text


def style_chart(chart, legend=True, size=11):
    chart.font.size = Pt(size)
    chart.font.name = FONT
    chart.font.color.rgb = rgb(MUTED)
    chart.has_title = False
    chart.has_legend = legend
    if legend:
        chart.legend.position = XL_LEGEND_POSITION.BOTTOM
        chart.legend.include_in_layout = False
        chart.legend.font.size = Pt(size)
        chart.legend.font.color.rgb = rgb(TEXT)
    va, ca = chart.value_axis, chart.category_axis
    va.major_gridlines.format.line.color.rgb = rgb(LINE)
    va.major_gridlines.format.line.width = Pt(0.5)
    va.format.line.fill.background()
    ca.format.line.color.rgb = rgb(LINE)
    va.tick_labels.font.color.rgb = rgb(MUTED)
    va.tick_labels.font.size = Pt(size)
    ca.tick_labels.font.color.rgb = rgb(TEXT)
    ca.tick_labels.font.size = Pt(size)


def color_series(chart, colors):
    for s, c in zip(chart.plots[0].series, colors):
        s.format.fill.solid()
        s.format.fill.fore_color.rgb = rgb(c)
        s.format.line.fill.background()


def labels(plot, fmt, size=11, pos=XL_LABEL_POSITION.OUTSIDE_END):
    plot.has_data_labels = True
    dl = plot.data_labels
    dl.number_format, dl.number_format_is_linked = fmt, False
    dl.font.size, dl.font.color.rgb, dl.font.name = Pt(size), rgb(TEXT), FONT
    dl.position = pos


def plain_table(tbl):
    pr = tbl._tbl.tblPr
    for ch_ in list(pr):
        if ch_.tag == qn("a:tableStyleId"):
            pr.remove(ch_)
    sid = etree.SubElement(pr, qn("a:tableStyleId"))
    sid.text = "{2D5ABB26-0587-4C30-8999-92F81FD0307C}"


def pct(v):  # JSON fraction -> percent number
    return round(v * 100, 3)


A = BENCH["ablations"]
HEAD = BENCH["headline"]
RAIL = BENCH["per_rail"]
HERO = BENCH["hero"]
SENS = BENCH["sensitivity"]

# =============================================================================================
# 1. Title
s = S("", None, data=False)
s.text(MX, 1.55, 6.4, 0.4, "DIGITAL PUBLIC SAFETY PLATFORM", 12, EMERALD, bold=True, count=False, spacing=200)
s.text(MX, 2.0, 6.4, 1.6, "ScamShield Immune", 44, TEXT, bold=True, count=False)
s.text(MX, 3.05, 6.2, 1.2, "One bank's detection protects every bank", 26, VIOLET, count=False)
s.text(MX, 4.7, 6.0, 0.6, "Reflex, shared memory and a fraud-network brain, with evidence a third party can verify.", 15, MUTED, count=False)
s.chip(MX, 5.7, 2.35, 0.4, "SIMULATED DATA", AMBER, 11)
s.text(MX + 2.55, 5.74, 3.5, 0.35, "Benchmark seed 42, hero run seed 1", 12, SUBTLE, count=False)
s.image("player-live-1440x900-dark.png", 7.1, 1.15, 5.65, 5.0)
s.notes(
    "Opening: one sentence. Digital-arrest scams move money in minutes and every bank sees only its own customers; "
    "this platform makes one bank's detection protect all of them, before the money moves.\n\n"
    "Say up front: everything in this talk is simulator data and real service code. The hero case is constructed.\n\n"
    "Sources: README.md (concept and quickstart), docs/architecture.md, docs/demo-script.md. Screenshot: docs/screenshots/player-live-1440x900-dark.png."
)

# 2. Problem
s = S("The problem", "Reactive investigation loses to scams that move in minutes")
s.stat(MX, 1.8, 5.9, 1.55, "1.14 million", "complaints, India, 2023", None, CRIMSON, 40)
s.stat(MX, 3.6, 5.9, 1.55, "Rs 1,776 crore", "digital-arrest scams, first nine months of 2024", None, AMBER, 40)
s.card(6.9, 1.8, 5.83, 3.35)
s.text(7.2, 2.0, 5.3, 0.4, "Why it fails", 16, TEXT, bold=True)
s.text(7.2, 2.6, 5.3, 2.4, [
    ("Money leaves before any report.", {"after": 10}),
    ("Banks see only their own customers.", {"after": 10}),
    ("Mules are reused on the next victim.", {"after": 10}),
], 15, MUTED)
s.text(MX, 5.6, 12, 0.6, "Source: challenge statement", 13, MUTED)
s.notes(
    "Talk track: the scale is the challenge statement's: 1.14 million cybercrime complaints in India in 2023, and digital-arrest scams above Rs 1,776 crore in the first nine months of 2024. "
    "Investigation after a complaint starts when the money is already gone; and a mule account that has already hurt a customer of bank A is invisible to bank B.\n\n"
    "Sources: both figures are quoted from the challenge statement as supplied with the brief; they are NOT in the repository's files, so they carry no repo citation. "
    "The framing sentence is from docs/superpowers/specs/2026-10-07-scam-immune-system-design.md section 1 (Intent)."
)

# 3. Insight
s = S("The insight", "An immune system for payments")
s.card(MX, 1.85, 12.13, 3.55, alpha=60)
# bank A -> hub -> banks B, C
s.card(1.0, 2.3, 2.5, 1.0, fill="123B30", line=EMERALD)
s.text(1.0, 2.52, 2.5, 0.5, "Bank A", 18, TEXT, bold=True, align="c", count=False)
s.text(1.0, 2.9, 2.5, 0.3, "detects and holds", 12, EMERALD, align="c")
s.arrow(3.5, 2.8, 5.0, 3.45, EMERALD, 2)
s.card(5.0, 3.0, 3.0, 1.0, fill="2A1F52", line=VIOLET)
s.text(5.0, 3.18, 3.0, 0.4, "Antibody hub", 18, TEXT, bold=True, align="c", count=False)
s.text(5.0, 3.58, 3.0, 0.3, "keyed hashes only", 12, VIOLET, align="c")
s.arrow(8.0, 3.35, 9.5, 2.8, VIOLET, 2)
s.arrow(8.0, 3.65, 9.5, 4.4, VIOLET, 2)
s.card(9.5, 2.3, 2.7, 1.0, fill="3A1A26", line=CRIMSON)
s.text(9.5, 2.52, 2.7, 0.5, "Bank B", 18, TEXT, bold=True, align="c", count=False)
s.text(9.5, 2.9, 2.7, 0.3, "blocks the same mule", 12, CRIMSON, align="c")
s.card(9.5, 3.9, 2.7, 1.0, fill="3A1A26", line=CRIMSON)
s.text(9.5, 4.12, 2.7, 0.5, "Bank C", 18, TEXT, bold=True, align="c", count=False)
s.text(9.5, 4.5, 2.7, 0.3, "blocks it too", 12, CRIMSON, align="c")
s.text(MX, 5.75, 12.1, 0.5, "Detect once, protect everywhere. Hold and verify before the money moves.", 18, TEXT, bold=True)
s.notes(
    "Talk track: biology metaphor kept literal. One detection becomes an antibody, a keyed hash of the mule account, applied at every bank. "
    "The hub never holds the key and never sees names or account numbers. Holds are verify steps, not blocks; they never auto-release or auto-block at the deadline.\n\n"
    "Sources: README.md (How it works), docs/architecture.md sections 3 and 4, services/antibody_hub/README.md (threat model and key handling), services/txn_guard/README.md (hold-and-verify)."
)

# 4. What we built
s = S("What we built", "Three layers, a citizen shield and evidence you can verify")
s.card(MX, 1.75, 2.55, 3.55, alpha=60)
s.text(MX + 0.2, 1.9, 2.2, 0.3, "SIGNALS (SIMULATED)", 11, SUBTLE, bold=True, count=False)
s.chip(MX + 0.2, 2.35, 2.15, 0.7, "Calls and messages", CYAN, 12)
s.chip(MX + 0.2, 3.3, 2.15, 0.7, "UPI / IMPS / NEFT", CYAN, 12)
s.arrow(3.2, 3.5, 3.55, 3.5, MUTED, 2)
layers = [
    ("REFLEX", "call-guard + txn-guard", "Hold and verify in minutes", EMERALD),
    ("MEMORY", "antibody-hub", "Shared keyed-hash memory", VIOLET),
    ("BRAIN", "graph-twin", "Rings and takedown what-if", AMBER),
]
for i, (nm, svc, d, c) in enumerate(layers):
    x = 3.6 + i * 2.1
    s.card(x, 1.75, 1.95, 3.55, line=c, lw=1.25)
    s.text(x + 0.15, 1.95, 1.65, 0.3, nm, 13, c, bold=True, count=False, spacing=100)
    s.text(x + 0.15, 2.5, 1.7, 0.9, svc.split(" + "), 15, TEXT, bold=True)
    s.text(x + 0.15, 3.5, 1.65, 1.6, d, 14, MUTED)
s.arrow(9.95, 3.5, 10.3, 3.5, MUTED, 2)
s.card(10.35, 1.75, 2.38, 3.55, alpha=60)
s.text(10.55, 1.9, 2.0, 0.3, "PEOPLE", 11, SUBTLE, bold=True, count=False)
s.chip(10.55, 2.35, 2.0, 0.55, "Citizen shield", EMERALD, 12)
s.chip(10.55, 3.05, 2.0, 0.55, "Analyst console", EMERALD, 12)
s.chip(10.55, 3.75, 2.0, 0.55, "Command centre", EMERALD, 12)
s.card(MX, 5.55, 12.13, 0.95, fill="0F2A3A", line=CYAN, lw=1.25)
s.text(MX + 0.25, 5.7, 5, 0.3, "EVIDENCE LEDGER", 11, CYAN, bold=True, count=False, spacing=100)
s.text(MX + 0.25, 6.0, 11.5, 0.4, "Hash chain, signed checkpoints, signed package, offline verify.py", 15, TEXT)
s.notes(
    "Talk track: reflex is per-transaction hold-and-verify; memory shares antibodies as keyed hashes; brain finds rings and simulates a takedown. "
    "The citizen shield scores a message in plain language in four languages. Everything the services decide is appended to a ledger that never sits on the hot path.\n\n"
    "Sources: README.md (architecture diagram and services table), docs/architecture.md sections 1 to 5, services/*/README.md. Counterfeit-currency detection was dropped by decision (spec section 1)."
)

# 5. Hero story
s = S("The hero run", "One constructed case, eight steps, 71 seconds at 1x")
s.image("player-live-1440x900-dark.png", MX, 1.6, 5.95, 3.7)
s.image("player-finale-1440x900-dark.png", 6.78, 1.6, 5.95, 3.7)
s.text(MX, 5.38, 5.95, 0.3, "Bank B blocks the same mule", 11, MUTED, count=False)
s.text(6.78, 5.38, 5.95, 0.3, "Victim B protected, evidence sealed", 11, MUTED, count=False)
steps = ["Call begins", "Call flagged", "Transfer held", "Analyst confirms", "Antibody shared", "Bank B blocks", "Cash-out", "Evidence sealed"]
cw, gap = 1.43, 0.1
for i, t in enumerate(steps):
    x = MX + i * (cw + gap)
    s.card(x, 5.85, cw, 0.85, fill=CARD, line=LINE)
    s.text(x + 0.1, 5.92, 0.4, 0.3, str(i + 1), 13, EMERALD, bold=True)
    s.text(x + 0.1, 6.2, cw - 0.2, 0.45, t, 11, TEXT)
s.notes(
    "Talk track (matches docs/demo-script.md, seed 1, cinematic pacing): victim A gets a scam call; call-guard flags it; her transfer to a new payee is held; "
    "an analyst confirms the mule; the antibody is shared; victim B at another bank pays the same mule 90 s later with an unflagged call and is held on ANTIBODY_MATCH; "
    "the ring cashes out on the map; evidence is sealed. Step offsets at 1x: held +24.7 s, confirmed and shared +27.2 s, bank B blocks +36.7 s, cash-out +51.2 s, sealed +71.1 s.\n\n"
    "Be explicit that this is a constructed hard case (seasoned mule, off-script call).\n\n"
    "Sources: docs/demo-script.md (Timings table), README.md (The story in eight steps), docs/screenshots/player-live-1440x900-dark.png and player-finale-1440x900-dark.png."
)

# 6. Under the hood
s = S("Under the hood", "Real services and models, a simulated world")
rows = [
    ("Call scoring", "clf-v1 classifier", "scripts, speakers"),
    ("Txn scoring", "gbm-v1 model", "accounts, transfers"),
    ("Threat memory", "hub store, outbox", "two banks"),
    ("Evidence", "Ed25519 chain, verifier", "simulated clock"),
    ("Analyst", "resolve API", "60 s oracle"),
]
tx, ty, tw = MX, 1.75, 12.13
tbl = s.s.shapes.add_table(len(rows) + 1, 3, Inches(tx), Inches(ty), Inches(tw), Inches(0.7 * (len(rows) + 1))).table
plain_table(tbl)
for ci, wv in enumerate([2.8, 5.0, 4.33]):
    tbl.columns[ci].width = Inches(wv)
for ri, row in enumerate([("Piece", "Real", "Simulated")] + rows):
    tbl.rows[ri].height = Inches(0.7)
    for ci, val in enumerate(row):
        c = tbl.cell(ri, ci)
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(CARD if ri else "1B2333")
        c.margin_left = Inches(0.18)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        p = c.text_frame.paragraphs[0]
        r = p.add_run()
        r.text = val
        r.font.name, r.font.size = FONT, Pt(16)
        r.font.bold = ri == 0 or ci == 0
        r.font.color.rgb = rgb(EMERALD if (ri == 0 and ci == 1) else AMBER if (ri == 0 and ci == 2) else TEXT)
        if ri:
            WORDS[s.n] += len(val.split())
s.text(MX, 6.25, 12.1, 0.5, "make demo runs the real services in one process.", 15, MUTED)
s.notes(
    "Talk track: the detection code is real and runs in-process; the world it sees is synthetic. No real telecom or bank integration exists. "
    "The demo is one process over an in-memory bus; the multi-container compose stack exists but was NOT run (no Docker daemon).\n\n"
    "Sources: web-console/src/lib/about/architecture.ts (REAL_VS_SIMULATED, shown on the console's /about page), README.md (Two packagings, What was and was not exercised), docs/architecture.md."
)

# 7. Results
s = S("Results", "Honest results: no false holds, with a weak NEFT rail")
hp, hr, hf = HEAD["precision"], HEAD["recall"], HEAD["fpr"]
s.stat(MX, 1.6, 2.9, 1.45, f"{pct(hp['value']):.1f}%", "Precision, hold", f"95% CI {pct(hp['ci95'][0]):.1f} to {pct(hp['ci95'][1]):.1f}", EMERALD, 30)
s.stat(3.65, 1.6, 2.9, 1.45, f"{pct(hr['value']):.1f}%", "Recall, scam txns", f"95% CI {pct(hr['ci95'][0]):.1f} to {pct(hr['ci95'][1]):.1f}", CYAN, 30)
s.stat(6.7, 1.6, 2.9, 1.45, f"{pct(hf['value']):.3f}%", "Benign held (FPR)", f"upper CI {pct(hf['ci95'][1]):.3f}%", EMERALD, 30)
s.stat(9.75, 1.6, 2.98, 1.45, "+6h 29m", "Mean lead time", "5 of 5 before 10th victim", AMBER, 30)
# per-rail chart
cd = CategoryChartData()
cd.categories = ["UPI", "IMPS", "NEFT"]
cd.add_series("Held (hold_verify)", [pct(RAIL[r]["fpr"]["value"]) for r in ("UPI", "IMPS", "NEFT")])
cd.add_series("Flagged (step-up or hold)", [pct(RAIL[r]["flagged_fpr"]["value"]) for r in ("UPI", "IMPS", "NEFT")])
s.card(MX, 3.25, 5.95, 3.55, alpha=60)
s.text(MX + 0.2, 3.33, 5.5, 0.3, "Benign false-positive rate by rail (%)", 13, TEXT, bold=True, count=False)
gf = s.s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(MX + 0.1), Inches(3.65), Inches(5.75), Inches(3.1), cd)
ch = gf.chart
style_chart(ch)
color_series(ch, [EMERALD, AMBER])
labels(ch.plots[0], '0.000"%"', 10)
ch.plots[0].gap_width = 70
ch.value_axis.maximum_scale = 1.0
ch.value_axis.minimum_scale = 0
ch.value_axis.tick_labels.number_format, ch.value_axis.tick_labels.number_format_is_linked = '0.0"%"', False
s.text(MX + 3.4, 3.38, 2.4, 0.3, "NEFT: no scam data", 11, AMBER, bold=True, align="r")
# ablation chart
order = [("baseline", "Baseline"), ("no_call_signal", "No call signal"), ("degraded_call_signal", "Degraded call"), ("no_antibody", "No antibody")]
cd = CategoryChartData()
cd.categories = [l for _, l in order]
cd.add_series("Recall, all scam txns", [pct(A[k]["recall"]["value"]) for k, _ in order])
cd.add_series("Victim-transfer recall", [pct(A[k]["victim_transfer_recall"]["value"]) for k, _ in order])
s.card(6.78, 3.25, 5.95, 3.55, alpha=60)
s.text(6.98, 3.33, 5.5, 0.3, "Ablations: recall (%), axis from 50", 13, TEXT, bold=True, count=False)
gf = s.s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(6.88), Inches(3.65), Inches(5.75), Inches(3.1), cd)
ch = gf.chart
style_chart(ch)
color_series(ch, [CYAN, VIOLET])
labels(ch.plots[0], '0.0', 10)
ch.plots[0].gap_width = 70
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 50, 110
ch.value_axis.major_unit = 10
ch.value_axis.tick_labels.number_format, ch.value_axis.tick_labels.number_format_is_linked = '0', False
s.notes(
    "Talk track: seed 42 simulator run, 21,736 benign and 680 scam transactions. Precision 100% [99.3, 100.0] because zero benign transactions were held "
    "(upper bound of the false-positive rate 0.018%); recall 82.6% [79.6, 85.3]. Precision depends on prevalence: the run is 3.03%, real prevalence is far lower (prevalence table in the report). "
    "Per rail: held benign is 0.000% everywhere, but flagged-benign (step-up or hold) is UPI 0.016%, IMPS 0.37%, NEFT 0.83%. NEFT is the worst rail and the simulator has no NEFT scams, "
    "so its figure is benign traffic only and its thresholds are hand-scaled. Ablation: removing antibodies drops victim-transfer recall from 100% to 94.5%; "
    "no call signal 98.1%; degraded call signal 99.2%. Lead time: 5 of 5 campaigns detected before the 10th victim, mean +6h29m13s at the first hold; few campaigns, wide uncertainty.\n\n"
    "All chart values are read from services/demo_server/demo_server/data/benchmark_summary.json at build time (headline, per_rail, ablations); "
    "lead time and prevalence from docs/benchmark_report.md. Simulator data, same generators as training, disjoint seeds.\n\n"
    "Note: the ablation axis starts at 50 so the differences are visible; say so out loud."
)

# 8. Hard-case honesty
s = S("Hard-case honesty", "Hard on purpose: antibodies must land early")
hm, hw = HERO["model_only"], HERO["with_antibody"]
cd = CategoryChartData()
cd.categories = ["Allow", "Step-up", "Hold-verify"]
cd.add_series("Model only", [hm["allow"], hm["step_up"], hm["hold_verify"]])
cd.add_series("With antibody", [hw["allow"], hw["step_up"], hw["hold_verify"]])
s.card(MX, 1.6, 5.95, 3.9, alpha=60)
s.text(MX + 0.2, 1.7, 5.5, 0.3, f"Victim B's first transfer, {HERO['seeds']} hero seeds", 13, TEXT, bold=True, count=False)
gf = s.s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(MX + 0.1), Inches(2.05), Inches(5.75), Inches(3.35), cd)
ch = gf.chart
style_chart(ch)
color_series(ch, [CRIMSON, EMERALD])
labels(ch.plots[0], '0', 11)
ch.plots[0].gap_width = 70
ch.value_axis.maximum_scale = 50
dl = [(SENS_i["label"], pct(SENS_i["victim_transfer_recall"]["value"])) for SENS_i in SENS]
short = {"no antibody": "None", "label-gated, 0 s": "0 s", "label-gated, 60 s (baseline)": "60 s", "label-gated, 300 s": "5 min",
         "label-gated, 900 s": "15 min", "label-gated, 1 h": "1 h", "label-gated, 6 h": "6 h", "label-gated, 24 h": "24 h"}
pts = [(short[l], v) for l, v in dl if l in short]
cd = CategoryChartData()
cd.categories = [p[0] for p in pts]
cd.add_series("Victim-transfer recall", [p[1] for p in pts])
s.card(6.78, 1.6, 5.95, 3.9, alpha=60)
s.text(6.98, 1.7, 5.5, 0.3, "Recall (%) vs analyst confirmation delay", 13, TEXT, bold=True, count=False)
gf = s.s.shapes.add_chart(XL_CHART_TYPE.LINE_MARKERS, Inches(6.88), Inches(2.05), Inches(5.75), Inches(3.35), cd)
ch = gf.chart
style_chart(ch, legend=False)
ser = ch.plots[0].series[0]
ser.format.line.color.rgb = rgb(VIOLET)
ser.format.line.width = Pt(2.5)
ser.marker.format.fill.solid()
ser.marker.format.fill.fore_color.rgb = rgb(VIOLET)
ser.marker.format.line.color.rgb = rgb(VIOLET)
ser.smooth = False
labels(ch.plots[0], '0.0', 10, XL_LABEL_POSITION.ABOVE)
ch.value_axis.minimum_scale, ch.value_axis.maximum_scale = 90, 104
ch.value_axis.major_unit = 2
ch.value_axis.tick_labels.number_format, ch.value_axis.tick_labels.number_format_is_linked = '0', False
s.text(MX, 5.75, 12.1, 0.9, [
    "Seasoned mule, off-script call: the model alone allowed 39 of 40; with the antibody all 40 were held.",
    "At a 24 h confirmation delay the benefit is gone (94.5%).",
], 15, TEXT, after=6)
s.notes(
    "Talk track: we built the hero to be the hardest case, not the typical one: a seasoned mule (68 to 398 days old, so no young-account signal) and an off-script call that call-guard misses. "
    "Across 40 hero seeds the model alone allowed victim B's first transfer 39 times and stepped it up once; with the antibody all 40 were held. The 40 seeds vary mule age, amounts and timing, not the script. "
    "Right chart: the confirmation delay is a free parameter. Victims in the simulator are minutes to hours apart, so delays up to 1 h cost nothing; at 6 h recall is 99.2%; at 24 h it is back to 94.5% and no antibody matches. "
    "This is the honest bound on 'detect once, protect everywhere': an antibody has to land before the next victim pays.\n\n"
    "Also say: antibody poisoning, false confirmations and merchant antibodies are not simulated; benign antibody matches = 0 holds by construction.\n\n"
    "Sources: benchmark_summary.json keys hero (model_only, with_antibody, seeds) and sensitivity; docs/benchmark_report.md (Antibody stage, Limitations); docs/judging-map.md."
)

# 9. Auditability
s = S("Auditability", "Evidence a third party can check without trusting us")
s.image("evidence-verified-1440x900-dark.png", MX, 1.6, 7.3, 4.6)
cards = [
    ("Proves", "Entries unaltered, in order, signed under a pinned key.", EMERALD),
    ("Does not prove", "That events were true, or legally admissible.", CRIMSON),
    ("Verify offline", "python -I -S verify.py --trusted-pubkey KEY", CYAN),
]
for i, (h, b, c) in enumerate(cards):
    y = 1.6 + i * 1.58
    s.card(8.25, y, 4.48, 1.42, line=c, lw=1.25)
    s.text(8.45, y + 0.15, 4.1, 0.3, h.upper(), 11, c, bold=True, count=False, spacing=100)
    s.text(8.45, y + 0.52, 4.1, 0.8, b, 13, TEXT, mono=(i == 2))
s.text(MX, 6.4, 12, 0.4, "Payloads hold ids, reason codes and amount buckets only; PII-shaped entries are quarantined.", 13, MUTED)
s.notes(
    "Talk track: every decision is appended to a hash-chained ledger with signed checkpoints; a case package ships its own verify.py, run isolated (python -I -S) and pinned to the ledger public key. "
    "The screenshot shows PASS_PINNED. Integrity is not truthfulness: it proves what the services recorded, in what order, unaltered. We make no admissibility claim; the Section 63 certificate is a template for counsel. "
    "A DB superuser can still edit rows; that is made detectable, not prevented. No external timestamping.\n\n"
    "Sources: docs/audit-trail.md, services/evidence_ledger/README.md, docs/architecture.md sections 5 and 8, docs/screenshots/evidence-verified-1440x900-dark.png, `make verify-package` (PASS, README 'What was exercised')."
)

# 10. Privacy and safety
s = S("Privacy and safety engineering", "Tokens and hashes only, and audit never blocks an alert")
items = [
    ("Tokens, not identities", "Keyed hashes only.", VIOLET),
    ("PII guard", "No amounts or free text.", CYAN),
    ("Never on the hot path", "Alerts publish before audit.", EMERALD),
    ("Fail safe", "Rules fallback; no auto-release.", AMBER),
]
for i, (h, b, c) in enumerate(items):
    x = MX + i * 3.06
    s.card(x, 1.7, 2.93, 2.6, line=c, lw=1.25)
    s.text(x + 0.2, 1.9, 2.55, 0.8, h, 18, TEXT, bold=True)
    s.text(x + 0.2, 3.0, 2.55, 1.2, b, 15, MUTED)
s.card(MX, 4.75, 12.13, 1.5, fill="0F2A3A", line=CYAN)
s.text(MX + 0.25, 4.95, 11.6, 0.3, "INDEPENDENT ADVERSARIAL REVIEW", 11, CYAN, bold=True, count=False, spacing=100)
s.text(MX + 0.25, 5.4, 11.6, 0.8, "Independent adversarial review of each service; about 15,000 ledger mutations, none passed.", 18, TEXT)
s.notes(
    "Talk track: privacy is structural, not a policy: the platform sees pseudonymous tokens and keyed hashes, the hub never holds the federation key, and the ledger refuses PII-shaped payloads. "
    "Audit is never awaited on the fraud path (outbox, circuit breaker, sweeper). Fail-safe: missing or tampered model artifacts fall back to rules; antibody bootstrap fails open (weaker protection, never a blocked payment); holds are flagged overdue, never auto-resolved.\n\n"
    "Review claim: README 'Development workflow' states one task at a time with an independent adversarial review after each and fix rounds until findings were closed. "
    "The about-15,000-mutation ledger attack campaign figure is from docs/judging-map.md (Technical excellence); the review ledger itself is a git-ignored working file under .superpowers/, so we give no per-service finding counts here. "
    "Test counts, if asked: about 1,490 Python tests plus 172 console unit tests, make test-all 13 of 13 (README).\n\n"
    "Sources: docs/architecture.md sections 4, 6 and 8; docs/audit-trail.md; docs/judging-map.md; README.md."
)

# 11. UX
s = S("Three audiences", "Command centre, analyst console and citizen shield")
s.image("command-1440x900-dark.png", MX, 1.6, 4.55, 2.85)
s.image("analyst-1440x900-dark.png", 5.35, 1.6, 4.55, 2.85)
s.image("citizen-verdict-ta-390x844-dark.png", 10.35, 1.45, 2.4, 3.2, crop_h=844)
s.text(MX, 4.6, 4.55, 0.3, "Command centre: map, ring, takedown", 11, MUTED, count=False)
s.text(5.35, 4.6, 4.55, 0.3, "Analyst: deadline-sorted hold queue", 11, MUTED, count=False)
s.text(10.35, 4.82, 2.4, 0.3, "Citizen verdict, Tamil", 11, MUTED, count=False)
s.chip(MX, 5.35, 3.9, 0.5, "English, Hindi, Tamil, Bengali", EMERALD, 12)
s.chip(4.7, 5.35, 3.9, 0.5, "axe-core: 0 serious or critical", CYAN, 12)
s.chip(8.8, 5.35, 3.93, 0.5, "Keyboard shortcuts, palette", VIOLET, 12)
s.text(MX, 6.1, 12.1, 0.6, "Risk is never shown by colour alone; one verdict per screen for citizens.", 14, MUTED)
s.notes(
    "Talk track: seven pages, three audiences. The officer sees the map, ring graph and a takedown what-if; the analyst works a deadline-sorted hold queue with countdown rings and a guarded resolve; "
    "the citizen pastes a message and gets one plain-language verdict in English, Hindi, Tamil or Bengali, nothing stored. 'Safe' only means no clear scam cues found.\n\n"
    "Accessibility: WCAG AA tokens; axe-core (wcag2a/aa, 2.1 a/aa) on all 7 pages in dark and light reported 0 serious or critical violations (web-console F2 report, a git-ignored working file). "
    "Automated rules only; the Playwright e2e and axe checks were not re-run in the final delivery pass (README 'What was and was not exercised'). Citizen reason lists stay in English in this build.\n\n"
    "Sources: README.md (UX), docs/judging-map.md (UX table), web-console/README.md, .superpowers/sdd/2026-10-07-scam-immune-system/web-console-f2-report.md; screenshots command-1440x900-dark, analyst-1440x900-dark, citizen-verdict-ta-390x844-dark (top 844 px)."
)

# 12. Scalability
s = S("Scalability and operations", "Designed to scale by partition key; not load-tested here")
rows = [
    ("txn.events", "payer_token"),
    ("call.events, call.risk", "victim_token"),
    ("txn.decisions", "txn_id"),
    ("antibody.published", "key_hash"),
]
tbl = s.s.shapes.add_table(len(rows) + 1, 2, Inches(MX), Inches(1.7), Inches(6.4), Inches(0.6 * (len(rows) + 1))).table
plain_table(tbl)
tbl.columns[0].width, tbl.columns[1].width = Inches(3.4), Inches(3.0)
for ri, row in enumerate([("Topic", "Partition key")] + rows):
    tbl.rows[ri].height = Inches(0.6)
    for ci, val in enumerate(row):
        c = tbl.cell(ri, ci)
        c.fill.solid()
        c.fill.fore_color.rgb = rgb(CARD if ri else "1B2333")
        c.margin_left = Inches(0.18)
        c.vertical_anchor = MSO_ANCHOR.MIDDLE
        r = c.text_frame.paragraphs[0].add_run()
        r.text = val
        r.font.name = "Courier New" if ri and ci == 0 else FONT
        r.font.size, r.font.bold = Pt(13), ri == 0
        r.font.color.rgb = rgb(EMERALD if ri == 0 else TEXT)
        if ri:
            WORDS[s.n] += len(val.split())
s.stat(7.4, 1.7, 2.65, 1.5, "7 ms", "p50 per transaction", "p99 about 14 ms", EMERALD, 30)
s.stat(10.2, 1.7, 2.53, 1.5, "3", "compose profiles", "core, stack, observability", CYAN, 30)
s.card(7.4, 3.45, 5.33, 1.2)
s.text(7.6, 3.75, 4.95, 0.8, "Audit outbox; stateless scorers.", 16, TEXT)
s.card(MX, 5.35, 12.13, 1.0, fill="3A2A0F", line=AMBER)
s.text(MX + 0.25, 5.65, 11.6, 0.6, "Latency is in-process. The Docker stack was validated, not run.", 16, TEXT)
s.notes(
    "Talk track: ordering is by payer: txn.events is keyed by payer_token and call.events/call.risk by victim_token so a call's risk meets that payer's transactions in order; antibody.published is keyed by key_hash so a revoke cannot overtake its create. "
    "Scorers are stateless apart from Redis and scale by consumer-group partitions; hub and ledger scale as replicas over one Postgres; graph-twin is a single replica (no sharding). "
    "Latency: consumer-path p50 about 7 ms and p99 about 14 ms per transaction with fakeredis stores (txn-guard README); the benchmark JSON's own latency_ms block reads p50 5.3 ms, p99 15.3 ms over 2,242 decisions. Target p99 under 300 ms. Not load-tested against real Kafka, Redis or Postgres.\n\n"
    "Sources: docs/architecture.md section 7 (Deployment notes) and section 2; contracts/scam_contracts/topics.py (via architecture.md); services/txn_guard/README.md; infra/docker-compose.yml (profiles); README.md (What was and was not exercised)."
)

# 13. Limits and next steps
s = S("Limits and next steps", "What we do not know yet, and what comes next")
s.card(MX, 1.7, 5.95, 3.9, alpha=60)
s.text(MX + 0.25, 1.85, 5.4, 0.35, "LIMITS", 12, AMBER, bold=True, count=False, spacing=100)
s.text(MX + 0.25, 2.35, 5.45, 3.8, [
    ("Simulator data only", {"after": 14}),
    ("No NEFT scam data", {"after": 14}),
    ("Modelled analyst", {"after": 14}),
    ("Stack not run", {"after": 14}),
], 20, TEXT)
s.card(6.78, 1.7, 5.95, 3.9, alpha=60)
s.text(7.03, 1.85, 5.4, 0.35, "NEXT STEPS", 12, EMERALD, bold=True, count=False, spacing=100)
s.text(7.03, 2.35, 5.45, 3.8, [
    ("1  Real data integration", {"after": 14}),
    ("2  NEFT scam data", {"after": 14}),
    ("3  Graph-twin on live traffic", {"after": 14}),
    ("4  Legal review of evidence", {"after": 14}),
    ("5  One-bank pilot", {"after": 14}),
], 20, TEXT)
s.notes(
    "Talk track: be direct about what the numbers are: a regression and ablation harness on simulator data, not real-world performance. NEFT has no scam transactions so its recall is undefined and its false-positive figure is benign traffic only. "
    "The analyst is an assumed oracle with a 60 s delay; poisoning is not simulated. The docker stack was validated with compose config only. "
    "Next steps are the pilot path: real data integration (call metadata, bank transaction feeds), NEFT scam data, graph-twin on live traffic (it is a single replica today), counsel review of the evidence package and the Section 63 certificate template, then one bank.\n\n"
    "Sources: README.md (Honest limitations), docs/benchmark_report.md (Limitations), docs/judging-map.md caveat columns, docs/architecture.md."
)

# 14. Close
s = S("", None, data=False)
s.text(MX, 1.5, 7.0, 0.4, "TRY IT", 12, EMERALD, bold=True, count=False, spacing=200)
s.text(MX, 1.95, 7.0, 1.2, "Run it in one command", 40, TEXT, bold=True, count=False)
s.card(MX, 3.35, 4.2, 0.9, fill=CARD, line=EMERALD, lw=1.25)
s.text(MX + 0.3, 3.62, 3.7, 0.5, "make demo", 24, EMERALD, bold=True, mono=True, count=False)
s.text(MX, 4.5, 7.0, 1.4, [
    "Open http://localhost:8080, choose Officer, press Play.",
    "github.com/aman-18-choudhary/ai_powered_immune (branch build/implementation)",
], 15, MUTED, count=False, after=8)
s.text(MX, 5.9, 6.0, 0.6, "Thank you", 28, VIOLET, bold=True, count=False)
s.chip(MX, 6.55, 2.35, 0.36, "SIMULATED DATA", AMBER, 10)
s.image("evidence-verified-1440x900-dark.png", 7.7, 1.5, 5.05, 3.2)
s.notes(
    "Close: one detection, every bank protected, every action auditable. Everything is simulated and says so. "
    "make demo builds the console and serves it at http://localhost:8080 (README quickstart). A recorded hero run with captions is in docs/demo/ (no narration; the script is docs/demo-script.md). "
    "The repository is https://github.com/aman-18-choudhary/ai_powered_immune; the work is on branch build/implementation, main holds only the design docs.\n\n"
    "Sources: README.md (60-second quickstart), docs/demo-script.md, docs/demo/README.md."
)

prs.save(OUT)
print("wrote", OUT)
for n, c in sorted(WORDS.items()):
    flag = "  <-- over 40" if c > 40 else ""
    print(f"slide {n:2d}: {c:3d} body words{flag}")
sys.exit(0)
